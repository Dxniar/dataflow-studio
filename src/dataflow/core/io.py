"""Local CSV/XLSX data IO with safe text exports and atomic, no-clobber publication."""

import codecs
import csv
import logging
import os
import tempfile
from collections.abc import Iterable
from pathlib import Path
from zipfile import ZipFile

import openpyxl
import pandas as pd
from openpyxl.cell import WriteOnlyCell

from .errors import DataFlowError
from .importing import ImportOptions, apply_import_schema
from .validation import validate_columns

logger = logging.getLogger(__name__)
SUPPORTED_EXTENSIONS = {".csv", ".xlsx"}


def _path(path: str | Path, *, must_exist: bool = False) -> Path:
    try:
        resolved = Path(path).expanduser().resolve()
        if resolved.suffix.lower() not in SUPPORTED_EXTENSIONS:
            raise DataFlowError("Unsupported file type. Choose a .csv or .xlsx file.")
        if must_exist and not resolved.is_file():
            raise DataFlowError(f"Input file does not exist or is not a file: {resolved.name}.")
        return resolved
    except (OSError, TypeError, RuntimeError):
        raise DataFlowError("Invalid or inaccessible file path.") from None


def validate_csv_options(encoding: str, delimiter: str) -> None:
    if not isinstance(encoding, str):
        raise DataFlowError("CSV encoding must be a supported encoding name.")
    try:
        codecs.lookup(encoding)
    except LookupError:
        raise DataFlowError("Unknown CSV encoding. Try utf-8-sig or cp1251.") from None
    if not isinstance(delimiter, str) or len(delimiter) != 1 or delimiter in '\r\n\x00"':
        raise DataFlowError("CSV delimiter must be one character, excluding quotes and newlines.")


def _checkpoint(control, stage=None, completed=0, total=0) -> None:
    if control is not None:
        if stage is not None:
            control.report(stage, completed, total)
        control.checkpoint()


def _guard_shape(rows: int, columns: int, options: ImportOptions) -> None:
    if rows > options.max_rows:
        raise DataFlowError("Import row limit exceeded. Reduce the source or increase max_rows.")
    if rows * columns > options.max_cells:
        raise DataFlowError("Import cell limit exceeded. Reduce the source or increase max_cells.")


def _guard_memory(frame: pd.DataFrame, options: ImportOptions) -> None:
    if frame.memory_usage(index=True, deep=True).sum() > options.max_table_mb * 1024 * 1024:
        raise DataFlowError(
            "Import table memory limit exceeded. Reduce the source or increase max_table_mb."
        )


def _guard_source(source: Path, options: ImportOptions) -> None:
    if source.stat().st_size > options.max_file_mb * 1024 * 1024:
        raise DataFlowError(
            "Import file size limit exceeded. Reduce the source or increase max_file_mb."
        )
    if source.suffix.lower() == ".xlsx":
        # Inspect ZIP metadata before openpyxl reads shared strings and workbook XML.
        with ZipFile(source) as archive:
            if (
                sum(item.file_size for item in archive.infolist())
                > options.max_table_mb * 1024 * 1024 * 4
            ):
                raise DataFlowError(
                    "Expanded workbook size limit exceeded. Reduce the workbook or increase max_table_mb."
                )


def _source_version(source: Path) -> tuple[int, int, int, int]:
    stat = source.stat()
    return stat.st_dev, stat.st_ino, stat.st_size, stat.st_mtime_ns


def _ensure_unchanged(source: Path, version: tuple[int, int, int, int]) -> None:
    try:
        unchanged = _source_version(source) == version
    except OSError:
        unchanged = False
    if not unchanged:
        raise DataFlowError("The source file changed during import. Close its editor and retry.")


def _validate_csv_structure(
    path: Path, encoding: str, delimiter: str, options: ImportOptions, control=None
) -> list[str]:
    with path.open(encoding=encoding, newline="") as stream:
        reader = csv.reader(stream, delimiter=delimiter, strict=True)
        try:
            header = next(reader)
        except StopIteration:
            raise DataFlowError("CSV file is empty. A header row is required.") from None
        validate_columns(header)
        text_bytes = 0
        for count, row in enumerate(reader, start=1):
            if count % 1000 == 1:
                _checkpoint(control, "read", count - 1, 0)
            if row and len(row) != len(header):
                raise DataFlowError(
                    "CSV rows have inconsistent field counts. Check its delimiter and quoting."
                )
            _guard_shape(count, len(header), options)
            text_bytes += sum(len(value.encode("utf-8")) for value in row)
            if text_bytes > options.max_table_mb * 1024 * 1024:
                raise DataFlowError(
                    "Import table memory limit exceeded. Reduce the source or increase max_table_mb."
                )
        return header


def _sheet_name(names: list[str], sheet: int | str) -> str:
    if type(sheet) is int and 0 <= sheet < len(names):
        return names[sheet]
    if isinstance(sheet, str) and sheet in names:
        return sheet
    raise DataFlowError(
        "Excel sheet was not found. Select an existing sheet name or zero-based index."
    )


def _validated_import_options(import_options: ImportOptions | dict | None) -> ImportOptions:
    return ImportOptions.from_dict(
        import_options.to_dict()
        if isinstance(import_options, ImportOptions)
        else {}
        if import_options is None
        else import_options
    )


def list_sheets(
    path: str | Path, *, import_options: ImportOptions | dict | None = None
) -> list[str]:
    source = _path(path, must_exist=True)
    if source.suffix.lower() != ".xlsx":
        raise DataFlowError("Sheet selection is only available for XLSX workbooks.")
    options = _validated_import_options(import_options)
    try:
        _guard_source(source, options)
        workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
        try:
            return list(workbook.sheetnames)
        finally:
            workbook.close()
    except DataFlowError:
        raise
    except Exception:
        # XLSX libraries can raise backend-specific XML errors containing raw data.
        raise DataFlowError(
            "Cannot read workbook. Check access and that it is a valid XLSX file."
        ) from None


def read_columns(
    path: str | Path,
    *,
    sheet: int | str = 0,
    encoding: str = "utf-8-sig",
    delimiter: str = ",",
    import_options: ImportOptions | dict | None = None,
) -> list[str]:
    """Read only the source header, without parsing or inferring any table cells."""
    source = _path(path, must_exist=True)
    options = _validated_import_options(import_options)
    try:
        _guard_source(source, options)
        if source.suffix.lower() == ".csv":
            validate_csv_options(encoding, delimiter)
            with source.open(encoding=encoding, newline="") as stream:
                header = next(csv.reader(stream, delimiter=delimiter, strict=True), [])
        else:
            workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
            try:
                name = _sheet_name(workbook.sheetnames, sheet)
                header = next(workbook[name].iter_rows(min_row=1, max_row=1, values_only=True), [])
            finally:
                workbook.close()
        return validate_columns(header)
    except DataFlowError:
        raise
    except Exception:
        raise DataFlowError(
            "Cannot read table header. Check the file, sheet, encoding, and delimiter."
        ) from None


def read_table(
    path: str | Path,
    *,
    sheet: int | str = 0,
    encoding: str = "utf-8-sig",
    delimiter: str = ",",
    import_options: ImportOptions | dict | None = None,
    control=None,
) -> pd.DataFrame:
    source = _path(path, must_exist=True)
    options = _validated_import_options(import_options)
    try:
        _checkpoint(control, "read", 0, 0)
        version = _source_version(source)
        _guard_source(source, options)
        if source.suffix.lower() == ".csv":
            validate_csv_options(encoding, delimiter)
            header = _validate_csv_structure(source, encoding, delimiter, options, control)
            if set(options.column_types) - set(header):
                raise DataFlowError(
                    "Import schema references a missing column. Check the source header."
                )
            _checkpoint(control, "parse", 0, 1)
            _ensure_unchanged(source, version)
            frame = pd.read_csv(
                source,
                encoding=encoding,
                sep=delimiter,
                keep_default_na=False,
                na_values=[""],
                skip_blank_lines=False,
                on_bad_lines="error",
                dtype={name: "string" for name in options.column_types},
                decimal=options.decimal,
                thousands=options.thousands,
                nrows=options.max_rows + 1,
            )
        else:
            with pd.ExcelFile(source, engine="openpyxl") as excel:
                name = _sheet_name(excel.sheet_names, sheet)
                header_rows = list(
                    excel.book[name].iter_rows(min_row=1, max_row=1, values_only=True)
                )
                if not header_rows:
                    raise DataFlowError("Excel sheet is empty. A header row is required.")
                validate_columns(header_rows[0])
                if set(options.column_types) - set(header_rows[0]):
                    raise DataFlowError(
                        "Import schema references a missing column. Check the source header."
                    )
                worksheet = excel.book[name]
                text_bytes = 0
                for count, row in enumerate(
                    worksheet.iter_rows(min_row=2, values_only=True), start=1
                ):
                    if count % 1000 == 1:
                        _checkpoint(control, "read", count - 1, 0)
                    _guard_shape(count, len(header_rows[0]), options)
                    text_bytes += sum(
                        len(value.encode("utf-8")) for value in row if isinstance(value, str)
                    )
                    if text_bytes > options.max_table_mb * 1024 * 1024:
                        raise DataFlowError(
                            "Import table memory limit exceeded. Reduce the source or increase max_table_mb."
                        )
                _checkpoint(control, "parse", 0, 1)
                _ensure_unchanged(source, version)
                frame = pd.read_excel(
                    excel,
                    sheet_name=name,
                    keep_default_na=False,
                    na_values=[""],
                    dtype={
                        column: "string" if kind == "string" else object
                        for column, kind in options.column_types.items()
                    },
                    decimal=options.decimal,
                    thousands=options.thousands,
                    nrows=options.max_rows + 1,
                )
        _ensure_unchanged(source, version)
        validate_columns(frame.columns)
        _guard_shape(len(frame), len(frame.columns), options)
        _guard_memory(frame, options)
        _checkpoint(control)
        frame = apply_import_schema(frame, options, control=control)
        _guard_memory(frame, options)
        _checkpoint(control, "read", len(frame), len(frame))
        _ensure_unchanged(source, version)
    except DataFlowError:
        raise
    except Exception:
        # Parsers often include complete cells/rows in their exceptions: never forward them.
        raise DataFlowError(
            f"Cannot read '{source.name}'. Check its format, encoding, delimiter, and file access."
        ) from None
    logger.info(
        "Read table: rows=%d columns=%d format=%s", len(frame), len(frame.columns), source.suffix
    )
    return frame


def _csv_text(value):
    if isinstance(value, str):
        stripped = value.lstrip(" \t\r\n\v\f\ufeff")
        if value.startswith(("\t", "\r", "\n")) or stripped.startswith(("=", "+", "-", "@")):
            return "'" + value
    return value


def _write_xlsx(frame: pd.DataFrame, temporary: Path, control=None) -> None:
    if len(frame) + 1 > 1_048_576 or len(frame.columns) > 16_384:
        raise DataFlowError("Table exceeds Excel's row or column limit. Export CSV instead.")
    workbook = openpyxl.Workbook(write_only=True)
    worksheet = workbook.create_sheet("Data")

    def cells(values):
        row = []
        for value in values:
            if pd.isna(value):
                value = None
            elif isinstance(value, pd.Timestamp):
                value = value.to_pydatetime()
            elif hasattr(value, "item"):
                value = value.item()
            cell = WriteOnlyCell(worksheet, value=value)
            if isinstance(value, str):
                if len(value) > 32_767:
                    raise DataFlowError(
                        "A text cell exceeds Excel's 32,767 character limit. Export CSV instead."
                    )
                cell.data_type = "s"
            row.append(cell)
        return row

    try:
        worksheet.append(cells(frame.columns))
        for number, values in enumerate(frame.itertuples(index=False, name=None)):
            if number % 1000 == 0:
                _checkpoint(control, "serialize", number, len(frame))
            worksheet.append(cells(values))
        _checkpoint(control, "serialize", len(frame), len(frame))
        workbook.save(temporary)
    finally:
        if not worksheet.closed:
            worksheet.close()
        workbook.close()


def write_table(
    frame: pd.DataFrame,
    path: str | Path,
    *,
    encoding: str = "utf-8-sig",
    delimiter: str = ",",
    collision: str = "error",
    source_paths: Iterable[str | Path] = (),
    control=None,
) -> Path:
    """Write completely to a sibling temp file, then atomically publish without overwrite.

    CSV formula-like strings receive an apostrophe; XLSX strings are explicitly text cells.
    Atomic publication uses a hard link on POSIX and non-replacing rename on Windows.
    """
    destination = _path(path)
    _checkpoint(control)
    if collision not in {"error", "unique"}:
        raise DataFlowError("Collision policy must be 'error' or 'unique'.")
    validate_columns(frame.columns)
    validate_csv_options(encoding, delimiter)
    for source in source_paths:
        try:
            if Path(source).expanduser().resolve() == destination:
                raise DataFlowError(
                    "Export cannot use an original input path. Choose a different filename."
                )
        except (OSError, TypeError, RuntimeError):
            raise DataFlowError("Cannot validate the original file path.") from None
    if destination.exists() and collision == "error":
        raise DataFlowError(f"Output already exists: {destination.name}. Choose another filename.")

    temporary = None
    try:
        destination.parent.mkdir(parents=True, exist_ok=True)
        handle, filename = tempfile.mkstemp(
            prefix=".dataflow-", suffix=destination.suffix, dir=destination.parent
        )
        os.close(handle)
        temporary = Path(filename)
        if destination.suffix.lower() == ".csv":
            with temporary.open("w", encoding=encoding, newline="") as stream:
                for start in range(0, max(1, len(frame)), 1000):
                    _checkpoint(control, "serialize", start, len(frame))
                    safe = frame.iloc[start : start + 1000].map(_csv_text)
                    safe.columns = [_csv_text(column) for column in frame.columns]
                    safe.to_csv(stream, index=False, sep=delimiter, header=start == 0)
        else:
            _write_xlsx(frame, temporary, control)
        # Windows FlushFileBuffers (used by fsync) requires a writable handle.
        with temporary.open("r+b") as stream:
            os.fsync(stream.fileno())
        candidate = destination
        number = 1
        while True:
            try:
                _checkpoint(control, "publish", 0, 1)
                if os.name == "nt":
                    # Windows rename fails if candidate exists; unlike POSIX rename, it cannot clobber.
                    os.rename(temporary, candidate)
                else:
                    os.link(temporary, candidate)
                break
            except FileExistsError:
                if collision == "error":
                    raise DataFlowError("Output already exists. Choose another filename.") from None
                candidate = destination.with_name(
                    f"{destination.stem}_{number}{destination.suffix}"
                )
                number += 1
        logger.info(
            "Exported table: rows=%d columns=%d format=%s",
            len(frame),
            len(frame.columns),
            candidate.suffix,
        )
        return candidate
    except DataFlowError:
        raise
    except Exception:
        raise DataFlowError(
            "Cannot export table. Check folder access, available space, encoding, and Excel cell limits."
        ) from None
    finally:
        if temporary is not None:
            temporary.unlink(missing_ok=True)
