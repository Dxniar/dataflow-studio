"""Real filesystem tests of Unicode, safe output, merge, and batch errors."""

import csv
from zipfile import ZipFile

import openpyxl
import pandas as pd
import pytest
from pandas.testing import assert_frame_equal

from dataflow.core.configuration import PipelineConfig
from dataflow.core.errors import DataFlowError
from dataflow.core.io import list_sheets, read_table, write_table
from dataflow.core.processor import batch_process, merge_files, process_file


def test_unicode_csv_excel_csv_conversion(tmp_path):
    frame = pd.DataFrame(
        {"name": ["Әлия", "Иван", "Jane"], "city": ["Қызылорда", "Москва", "York"]}
    )
    csv_path = write_table(frame, tmp_path / "исходный.csv", delimiter=";")
    loaded = read_table(csv_path, delimiter=";")
    xlsx_path = write_table(loaded, tmp_path / "кесте.xlsx")
    assert list_sheets(xlsx_path) == ["Data"]
    from_excel = read_table(xlsx_path, sheet="Data")
    output = write_table(from_excel, tmp_path / "result.csv")
    assert_frame_equal(read_table(output), frame)


def test_named_excel_sheet_is_respected(tmp_path):
    path = tmp_path / "multi.xlsx"
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        pd.DataFrame({"id": [1]}).to_excel(writer, sheet_name="First", index=False)
        pd.DataFrame({"id": [2]}).to_excel(writer, sheet_name="Second", index=False)
    assert list_sheets(path) == ["First", "Second"]
    assert read_table(path, sheet="Second")["id"].tolist() == [2]
    with pytest.raises(DataFlowError):
        read_table(path, sheet="Absent")


@pytest.mark.parametrize(
    "name,data",
    [
        ("empty.csv", b""),
        ("corrupt.xlsx", b"not a workbook"),
        ("duplicate.csv", b"id,id\n1,2\n"),
        ("bad.csv", b'id,name\n1,"unclosed\n'),
        ("unknown.txt", b"a,b\n1,2\n"),
    ],
)
def test_bad_inputs_have_domain_errors(tmp_path, name, data):
    path = tmp_path / name
    path.write_bytes(data)
    with pytest.raises(DataFlowError):
        read_table(path)


def test_missing_input_and_invalid_csv_settings(tmp_path):
    with pytest.raises(DataFlowError):
        read_table(tmp_path / "absent.csv")
    path = tmp_path / "data.csv"
    path.write_text("id\n1\n", encoding="utf-8")
    with pytest.raises(DataFlowError):
        read_table(path, encoding="not-an-encoding")
    with pytest.raises(DataFlowError):
        read_table(path, delimiter="||")


def test_malformed_workbook_xml_has_a_domain_error_for_sheet_listing(tmp_path):
    original = tmp_path / "original.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.append(["id"])
    workbook.save(original)
    workbook.close()
    damaged = tmp_path / "damaged.xlsx"
    with ZipFile(original) as source, ZipFile(damaged, "w") as target:
        for name in source.namelist():
            target.writestr(name, b"<broken" if name == "xl/workbook.xml" else source.read(name))
    with pytest.raises(DataFlowError, match="Cannot read workbook"):
        list_sheets(damaged)


def test_existing_files_and_originals_never_overwritten(tmp_path):
    frame = pd.DataFrame({"id": [1]})
    path = write_table(frame, tmp_path / "data.csv")
    original_bytes = path.read_bytes()
    with pytest.raises(DataFlowError):
        write_table(pd.DataFrame({"id": [2]}), path)
    assert path.read_bytes() == original_bytes
    with pytest.raises(DataFlowError):
        write_table(frame, path, collision="unique", source_paths=[path])
    unique = write_table(frame, path, collision="unique")
    assert unique != path
    assert path.read_bytes() == original_bytes
    assert_frame_equal(read_table(unique), frame)


@pytest.mark.parametrize("suffix", ["csv", "xlsx"])
def test_formula_looking_text_is_not_executable(tmp_path, suffix):
    values = ["=2+2", "+cmd", "-cmd", "@SUM(A1)", "  =1", "\t=1", "normal"]
    frame = pd.DataFrame({"=header": values, "number": [-1] * len(values)})
    output = write_table(frame, tmp_path / f"safe.{suffix}")
    if suffix == "xlsx":
        workbook = openpyxl.load_workbook(output, data_only=False)
        try:
            sheet = workbook.active
            assert all(cell.data_type != "f" for row in sheet for cell in row)
            assert sheet["A2"].value == "=2+2"
            assert sheet["B2"].value == -1
        finally:
            workbook.close()
    else:
        with output.open(encoding="utf-8-sig", newline="") as stream:
            rows = list(csv.reader(stream))
        assert rows[0][0] == "'=header"
        assert all(row[0].startswith("'") for row in rows[1:-1])
        assert rows[-1][0] == "normal"
        assert rows[1][1] == "-1"


def test_write_failure_leaves_no_output_or_temporary_file(tmp_path, monkeypatch):
    frame = pd.DataFrame({"id": [1]})

    def fail_write(*args, **kwargs):
        raise PermissionError("private-cell")

    monkeypatch.setattr(pd.DataFrame, "to_csv", fail_write)
    with pytest.raises(DataFlowError) as error:
        write_table(frame, tmp_path / "failed.csv")
    assert "private-cell" not in str(error.value)
    assert list(tmp_path.iterdir()) == []


def test_header_only_table_is_valid(tmp_path):
    path = tmp_path / "header.csv"
    path.write_text("id,name\n", encoding="utf-8")
    assert read_table(path).shape == (0, 2)


def test_csv_reader_preserves_blank_rows_and_literal_na(tmp_path):
    path = tmp_path / "literal.csv"
    path.write_text("id,city\n1,NA\n,\n2,NULL\n", encoding="utf-8")
    frame = read_table(path)
    assert len(frame) == 3
    assert frame["city"].iloc[0] == "NA"
    assert frame["city"].iloc[2] == "NULL"
    assert frame.iloc[1].isna().all()


def test_duplicate_excel_headers_are_rejected(tmp_path):
    path = tmp_path / "duplicate.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.append(["id", "id"])
    workbook.active.append([1, 2])
    workbook.save(path)
    workbook.close()
    with pytest.raises(DataFlowError):
        read_table(path)


def test_xlsx_formula_cells_are_not_evaluated(tmp_path):
    path = tmp_path / "formula.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.append(["id", "calculated"])
    workbook.active.append([1, "=1+1"])
    workbook.save(path)
    workbook.close()
    frame = read_table(path)
    assert pd.isna(frame.loc[0, "calculated"])


def test_output_directory_failure_is_a_domain_error(tmp_path):
    parent = tmp_path / "file.csv"
    parent.write_text("unchanged", encoding="utf-8")
    with pytest.raises(DataFlowError):
        write_table(pd.DataFrame({"id": [1]}), parent / "result.csv")
    assert parent.read_text(encoding="utf-8") == "unchanged"


def test_merge_union_intersection_and_strict_modes(tmp_path):
    first = write_table(pd.DataFrame({"id": [1], "name": ["A"]}), tmp_path / "a.csv")
    second = write_table(pd.DataFrame({"id": [2], "city": ["B"]}), tmp_path / "b.csv")
    union = merge_files([first, second])
    assert union.processed.columns.tolist() == ["id", "name", "city"]
    assert union.processed["id"].tolist() == [1, 2]
    assert union.summary()["files_processed"] == 2
    intersection = merge_files([first, second], mode="intersection")
    assert intersection.processed.columns.tolist() == ["id"]
    with pytest.raises(DataFlowError):
        merge_files([first, second], mode="strict")


def test_merge_reorders_matching_columns_and_rejects_incompatible_types(tmp_path):
    first = write_table(pd.DataFrame({"id": [1], "name": ["A"]}), tmp_path / "a.csv")
    second = write_table(pd.DataFrame({"name": ["B"], "id": [2]}), tmp_path / "b.csv")
    assert merge_files([first, second], mode="strict").processed["id"].tolist() == [1, 2]
    bad = write_table(pd.DataFrame({"id": ["private-cell"], "name": ["C"]}), tmp_path / "bad.csv")
    with pytest.raises(DataFlowError) as error:
        merge_files([first, bad])
    assert "private-cell" not in str(error.value)


def test_batch_continues_after_error_and_uses_unique_names(tmp_path):
    first = write_table(pd.DataFrame({"id": [1, 1]}), tmp_path / "first.csv")
    second = write_table(pd.DataFrame({"id": [2, 2]}), tmp_path / "second.csv")
    outputs = tmp_path / "outputs"
    result = batch_process(
        [first, tmp_path / "absent.csv", second], PipelineConfig(drop_duplicates=True), outputs
    )
    assert len(result.successes) == 2
    assert len(result.errors) == 1
    assert result.summary()["processed"] == 2
    assert result.summary()["failed"] == 1
    assert result.successes[0]["summary"]["duplicates_removed"] == 1
    assert len(read_table(first)) == 2
    again = batch_process([first], PipelineConfig(), outputs)
    assert again.successes[0]["output"] != result.successes[0]["output"]


def test_process_file_records_sources_and_metrics(tmp_path):
    path = write_table(pd.DataFrame({"id": [1, 1]}), tmp_path / "in.csv")
    result = process_file(path, PipelineConfig(drop_duplicates=True))
    assert result.source_paths == [str(path.resolve())]
    assert result.summary()["duplicates_removed"] == 1
